"""
build_index.py
Lee los documentos de /documents y construye el índice vectorial con ChromaDB.
Se ejecuta al inicio de cada sesión via setup.sh

Formatos soportados:
- Texto plano: .txt, .md, .csv, .json, .yaml, .yml, .html, .xml, .rst
- Código fuente: .py, .js, .ts, .java, .cpp, .c, .go, .rb, .sh, etc.
- Documentos binarios: .pdf, .docx, .xlsx
- Cualquier otro archivo se intenta leer como UTF-8; se omite si es binario.
"""

import os
import chromadb
from sentence_transformers import SentenceTransformer
from pypdf import PdfReader

DOCS_DIR = os.path.join(os.path.dirname(__file__), "../documents")
CHROMA_DIR = os.path.join(os.path.dirname(__file__), "../chroma_db")

# Extensiones que sabemos que son texto plano
TEXT_EXTENSIONS = {
    ".txt", ".md", ".markdown", ".csv", ".json", ".yaml", ".yml",
    ".html", ".htm", ".xml", ".rst", ".toml", ".ini", ".cfg", ".env",
    ".py", ".js", ".ts", ".jsx", ".tsx", ".java", ".cpp", ".c", ".h",
    ".go", ".rb", ".sh", ".bash", ".zsh", ".sql", ".r", ".scala",
    ".kt", ".swift", ".cs", ".php", ".rs",
}


def _read_as_text(filepath):
    """Intenta leer el archivo como texto UTF-8; retorna None si falla."""
    try:
        with open(filepath, "r", encoding="utf-8") as f:
            return f.read()
    except (UnicodeDecodeError, PermissionError):
        return None


def _read_pdf(filepath):
    reader = PdfReader(filepath)
    return "\n".join(
        page.extract_text() for page in reader.pages if page.extract_text()
    )


def _read_docx(filepath):
    try:
        import docx
        doc = docx.Document(filepath)
        return "\n".join(p.text for p in doc.paragraphs if p.text)
    except ImportError:
        print(f"  [!] python-docx no instalado, omitiendo {os.path.basename(filepath)}")
        return None


def _read_xlsx(filepath):
    try:
        import openpyxl
        wb = openpyxl.load_workbook(filepath, read_only=True, data_only=True)
        lines = []
        for sheet in wb.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_text = "\t".join(str(c) for c in row if c is not None)
                if row_text:
                    lines.append(row_text)
        return "\n".join(lines)
    except ImportError:
        print(f"  [!] openpyxl no instalado, omitiendo {os.path.basename(filepath)}")
        return None


def load_documents():
    docs = []
    for filename in sorted(os.listdir(DOCS_DIR)):
        filepath = os.path.join(DOCS_DIR, filename)
        if not os.path.isfile(filepath):
            continue

        ext = os.path.splitext(filename)[1].lower()
        text = None

        if ext == ".pdf":
            text = _read_pdf(filepath)
        elif ext == ".docx":
            text = _read_docx(filepath)
        elif ext == ".xlsx":
            text = _read_xlsx(filepath)
        elif ext in TEXT_EXTENSIONS:
            text = _read_as_text(filepath)
        else:
            # Intentar como texto genérico
            text = _read_as_text(filepath)
            if text is None:
                print(f"  [~] Omitido (binario no soportado): {filename}")
                continue

        if text and text.strip():
            docs.append({"filename": filename, "text": text})
            print(f"  [+] Cargado: {filename}")
        else:
            print(f"  [~] Omitido (vacío): {filename}")

    return docs


def chunk_text(text, chunk_size=500, overlap=50):
    words = text.split()
    chunks = []
    for i in range(0, len(words), chunk_size - overlap):
        chunk = " ".join(words[i:i + chunk_size])
        if chunk:
            chunks.append(chunk)
    return chunks


def build_index():
    print(f"Buscando documentos en {DOCS_DIR} ...")
    docs = load_documents()
    if not docs:
        print("No hay documentos en /documents. Agrega archivos y vuelve a ejecutar.")
        return

    print(f"\nCargando modelo de embeddings...")
    model = SentenceTransformer("all-MiniLM-L6-v2")

    client = chromadb.PersistentClient(path=CHROMA_DIR)
    try:
        client.delete_collection("documents")
    except Exception:
        pass
    collection = client.create_collection("documents")

    all_chunks, all_ids, all_metadata = [], [], []
    for doc in docs:
        chunks = chunk_text(doc["text"])
        for i, chunk in enumerate(chunks):
            all_chunks.append(chunk)
            all_ids.append(f"{doc['filename']}_chunk_{i}")
            all_metadata.append({"source": doc["filename"]})

    print(f"Generando embeddings para {len(all_chunks)} fragmentos...")
    embeddings = model.encode(all_chunks).tolist()

    collection.add(
        documents=all_chunks,
        embeddings=embeddings,
        ids=all_ids,
        metadatas=all_metadata,
    )

    print(f"\nÍndice construido: {len(all_chunks)} fragmentos de {len(docs)} documento(s).")


if __name__ == "__main__":
    build_index()
